#!/usr/bin/env python3
"""Summarize tool behavior and failure counts from a prompted-tool VLMEval run."""

from __future__ import annotations

import argparse
from collections import Counter
import json
from pathlib import Path
from typing import Any, Iterable


EXPECTED_TOOLS = {"crop_zoom", "grounding_detect", "object_count", "ocr_read", "ground_depth"}


def _is_api_failure(value: Any) -> bool:
    text = "" if value is None else str(value).strip().casefold()
    return any(
        marker in text
        for marker in ("api failed", "api failure", "failed to obtain answer via api")
    )


def _is_empty(value: Any) -> bool:
    return value is None or not str(value).strip() or str(value).strip().casefold() == "nan"


def _successful(call: dict[str, Any]) -> bool:
    if call.get("error"):
        return False
    result = call.get("result")
    if isinstance(result, dict) and str(result.get("status", "")).casefold() in {"error", "failed"}:
        return False
    return result is not None


def _raw_rows(path: Path) -> Iterable[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to summarize VLMEval .xlsx files") from exc

    worksheet = load_workbook(path, read_only=True, data_only=True).active
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = [str(value) if value is not None else "" for value in next(rows)]
    except StopIteration:
        return
    if "raw_response" not in headers:
        return
    for values in rows:
        yield dict(zip(headers, values))


def summarize_workbook(path: Path) -> dict[str, Any]:
    calls: Counter[str] = Counter()
    successes: Counter[str] = Counter()
    samples_with_tool: Counter[str] = Counter()
    total = api_failures = empty_predictions = parse_errors = tool_samples = turns = 0

    for row in _raw_rows(path):
        total += 1
        prediction = row.get("prediction")
        api_failures += int(_is_api_failure(prediction))
        empty_predictions += int(_is_empty(prediction))
        try:
            trace = json.loads(str(row.get("raw_response") or "{}"))
        except (TypeError, ValueError):
            parse_errors += 1
            continue
        turns += int(trace.get("turns") or 0)
        row_calls = trace.get("tool_calls") or []
        if not isinstance(row_calls, list):
            parse_errors += 1
            continue
        if row_calls:
            tool_samples += 1
        row_names = set()
        for call in row_calls:
            if not isinstance(call, dict):
                parse_errors += 1
                continue
            name = str(call.get("name") or "<missing>")
            calls[name] += 1
            row_names.add(name)
            if _successful(call):
                successes[name] += 1
        samples_with_tool.update(row_names)

    known_calls = sum(calls[name] for name in EXPECTED_TOOLS)
    known_successes = sum(successes[name] for name in EXPECTED_TOOLS)
    return {
        "workbook": str(path),
        "samples": total,
        "api_failures": api_failures,
        "empty_predictions": empty_predictions,
        "raw_response_parse_errors": parse_errors,
        "samples_using_any_tool": tool_samples,
        "tool_use_rate": round(tool_samples / total, 6) if total else 0.0,
        "mean_turns": round(turns / total, 6) if total else 0.0,
        "mean_tool_calls": round(sum(calls.values()) / total, 6) if total else 0.0,
        "expected_tool_calls": dict(sorted((name, calls[name]) for name in EXPECTED_TOOLS)),
        "expected_tool_sample_counts": dict(
            sorted((name, samples_with_tool[name]) for name in EXPECTED_TOOLS)
        ),
        "expected_tool_successes": dict(
            sorted((name, successes[name]) for name in EXPECTED_TOOLS)
        ),
        "expected_tool_success_rate": round(known_successes / known_calls, 6) if known_calls else 0.0,
        "unexpected_tool_calls": dict(
            sorted((name, count) for name, count in calls.items() if name not in EXPECTED_TOOLS)
        ),
    }


def aggregate_modes(workbooks: list[dict[str, Any]]) -> dict[str, Any]:
    aggregated: dict[str, dict[str, Any]] = {}
    for item in workbooks:
        mode = item["mode"]
        current = aggregated.setdefault(
            mode,
            {
                "datasets": [],
                "samples": 0,
                "api_failures": 0,
                "empty_predictions": 0,
                "raw_response_parse_errors": 0,
                "samples_using_any_tool": 0,
                "turn_sum": 0.0,
                "call_sum": 0.0,
                "expected_tool_calls": Counter(),
                "expected_tool_sample_counts": Counter(),
                "expected_tool_successes": Counter(),
                "unexpected_tool_calls": Counter(),
            },
        )
        samples = item["samples"]
        current["datasets"].append(item["dataset"])
        for key in (
            "samples",
            "api_failures",
            "empty_predictions",
            "raw_response_parse_errors",
            "samples_using_any_tool",
        ):
            current[key] += item[key]
        current["turn_sum"] += item["mean_turns"] * samples
        current["call_sum"] += item["mean_tool_calls"] * samples
        for key in (
            "expected_tool_calls",
            "expected_tool_sample_counts",
            "expected_tool_successes",
            "unexpected_tool_calls",
        ):
            current[key].update(item[key])

    result = {}
    for mode, current in aggregated.items():
        samples = current.pop("samples")
        calls = current["expected_tool_calls"]
        successes = current["expected_tool_successes"]
        current["samples"] = samples
        current["tool_use_rate"] = round(current["samples_using_any_tool"] / samples, 6) if samples else 0.0
        current["mean_turns"] = round(current.pop("turn_sum") / samples, 6) if samples else 0.0
        current["mean_tool_calls"] = round(current.pop("call_sum") / samples, 6) if samples else 0.0
        current["expected_tool_success_rate"] = (
            round(sum(successes.values()) / sum(calls.values()), 6) if sum(calls.values()) else 0.0
        )
        for key in (
            "expected_tool_calls",
            "expected_tool_sample_counts",
            "expected_tool_successes",
            "unexpected_tool_calls",
        ):
            current[key] = dict(sorted(current[key].items()))
        result[mode] = current
    return result


def summarize(root: Path) -> dict[str, Any]:
    root = root.resolve()
    workbooks = []
    for path in sorted(root.rglob("VisualAgent-vllm_*.xlsx")):
        if not path.is_file() or "_exact_matching_result" in path.name:
            continue
        relative = path.relative_to(root)
        mode = relative.parts[0] if relative.parts else "unknown"
        dataset = path.stem.removeprefix("VisualAgent-vllm_")
        item = summarize_workbook(path)
        if item["samples"] == 0:
            continue
        item.update({"mode": mode, "dataset": dataset})
        workbooks.append(item)

    score_artifacts = sorted(
        str(path.relative_to(root))
        for pattern in ("*_acc.csv", "*_score.json", "*_rating.json")
        for path in root.rglob(pattern)
        if path.is_file()
    )
    return {
        "root": str(root),
        "expected_tools": sorted(EXPECTED_TOOLS),
        "mode_totals": aggregate_modes(workbooks),
        "workbooks": workbooks,
        "score_artifacts": score_artifacts,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("root", type=Path, help="Experiment group root containing tool-evaluation outputs")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    result = summarize(args.root)
    serialized = json.dumps(result, ensure_ascii=False, indent=2) + "\n"
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(serialized, encoding="utf-8")
    print(serialized, end="")


if __name__ == "__main__":
    main()
